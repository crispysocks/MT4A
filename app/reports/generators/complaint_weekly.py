"""投诉处理周报生成器"""
from datetime import datetime
from typing import Optional
import os
import json

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlmodel import Session, select
from anthropic import Anthropic

from app.reports.base import ReportGenerator
from app.db.models import Complaint
from app.db.connection import engine


class ComplaintWeeklyGenerator(ReportGenerator):
    """投诉处理周报生成器"""
    
    REPORT_NAME = "投诉处理周报"
    
    def __init__(self, period: str = "this_week", start_date: Optional[datetime] = None, end_date: Optional[datetime] = None):
        super().__init__(period, start_date, end_date)
        self.llm = Anthropic(
            base_url=os.getenv("LLM_BASE_URL"),
            api_key=os.getenv("LLM_AUTH_TOKEN"),
        )
        self.model = os.environ.get("MODEL_ID", "qwen3.6-plus")
    
    def _get_filename_prefix(self) -> str:
        from app.reports.utils import format_period_label
        period_label = format_period_label(self.period)
        date_str = self.generated_at.strftime("%Y%m%d")
        return f"投诉处理周报_{period_label}_{date_str}"
    
    def collect_data(self) -> dict:
        """收集投诉数据"""
        start_date = self.start_date
        end_date = self.end_date
        
        if start_date is None or end_date is None:
            from app.reports.utils import get_date_range
            start_date, end_date = get_date_range(self.period)
        
        with Session(engine) as session:
            # 投诉记录
            complaints = session.exec(
                select(Complaint).where(
                    Complaint.created_at >= start_date,
                    Complaint.created_at <= end_date
                )
            ).all()
            
            # 按状态统计
            status_counts = {}
            for c in complaints:
                status = c.status or "未知"
                status_counts[status] = status_counts.get(status, 0) + 1
            
            # 按分类统计
            category_counts = {}
            for c in complaints:
                cat = c.category or "未分类"
                category_counts[cat] = category_counts.get(cat, 0) + 1
            
            # 已解决和未解决
            resolved = sum(1 for c in complaints if c.status == "resolved")
            pending = sum(1 for c in complaints if c.status in ("pending", "processing"))
            
            return {
                "complaints": complaints,
                "status_counts": status_counts,
                "category_counts": category_counts,
                "total": len(complaints),
                "resolved": resolved,
                "pending": pending,
                "start_date": start_date,
                "end_date": end_date,
            }
    
    def _ai_analyze(self, data: dict) -> str:
        """使用 LLM 生成投诉分析"""
        prompt = f"""作为售后服务分析专家，请根据以下投诉数据生成周报：

- 投诉总量：{data['total']} 件
- 状态分布：{json.dumps(data['status_counts'], ensure_ascii=False)}
- 分类分布：{json.dumps(data['category_counts'], ensure_ascii=False)}
- 已解决：{data['resolved']} 件 | 待处理：{data['pending']} 件

请分析：
1. 投诉总量及趋势分析
2. 智能分类统计（主要投诉类型）
3. 处理时效评估
4. 长期未决疑难案件预警
5. 改进建议

请直接输出分析内容。"""
        
        try:
            response = self.llm.messages.create(
                model=self.model,
                system="你是一个专业的售后服务分析专家。",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=2000,
            )
            
            for block in response.content:
                if hasattr(block, "text") and block.text:
                    return block.text.strip()
        except Exception as e:
            print(f"[AI 分析失败] {e}")
        
        return "AI 分析暂时不可用。"
    
    def build_content(self, doc: Document, data: dict) -> None:
        """构建报告内容"""
        # 1. 总体概况
        doc.add_heading("一、总体概况", level=1)
        overview = doc.add_paragraph()
        overview.add_run(f"投诉总量：").bold = True
        overview.add_run(f"{data['total']} 件\n")
        overview.add_run(f"已解决：").bold = True
        overview.add_run(f"{data['resolved']} 件\n")
        overview.add_run(f"待处理：").bold = True
        overview.add_run(f"{data['pending']} 件")
        
        # 2. 分类统计
        doc.add_heading("二、投诉分类统计", level=1)
        for category, count in data['category_counts'].items():
            doc.add_paragraph(f"{category}：{count} 件")
        
        # 3. 状态分布
        doc.add_heading("三、处理状态分布", level=1)
        for status, count in data['status_counts'].items():
            doc.add_paragraph(f"{status}：{count} 件")
        
        # 4. 待处理投诉明细
        if data['pending'] > 0:
            doc.add_heading("四、待处理投诉明细", level=1)
            for c in data['complaints']:
                if c.status in ("pending", "processing"):
                    para = doc.add_paragraph()
                    para.add_run(f"学生ID：{c.student_id} | 分类：{c.category or '未分类'} | 状态：{c.status}")
                    para.add_run(f"\n内容：{c.content[:200]}...")
        
        # 5. AI 分析
        doc.add_heading("五、AI 投诉分析", level=1)
        analysis = self._ai_analyze(data)
        doc.add_paragraph(analysis)
