"""报告生成器基类"""
from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ReportGenerator(ABC):
    """报告生成器抽象基类"""
    
    REPORT_NAME: str = "未命名报告"
    REPORT_DIR: Path = Path("reports")
    
    def __init__(self, period: str = "this_week", start_date: Optional[datetime] = None, end_date: Optional[datetime] = None):
        """
        Args:
            period: 周期标识 (today/this_week/last_week 等)
            start_date: 自定义开始日期（可选）
            end_date: 自定义结束日期（可选）
        """
        self.period = period
        self.start_date = start_date
        self.end_date = end_date
        self.generated_at = datetime.now()
    
    @abstractmethod
    def collect_data(self) -> dict:
        """收集报告所需数据，返回字典格式"""
        pass
    
    @abstractmethod
    def build_content(self, doc: Document, data: dict) -> None:
        """构建报告内容，向 Document 对象添加内容"""
        pass
    
    def generate(self) -> Path:
        """
        生成报告文件
        
        Returns:
            生成的文件路径
        """
        self.REPORT_DIR.mkdir(parents=True, exist_ok=True)
        
        # 生成文件名
        timestamp = self.generated_at.strftime("%Y%m%d_%H%M%S")
        filename = f"{self._get_filename_prefix()}_{timestamp}.docx"
        filepath = self.REPORT_DIR / filename
        
        # 创建文档
        doc = Document()
        self._setup_styles(doc)
        
        # 添加标题
        title = doc.add_heading(self.REPORT_NAME, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 添加元信息
        self._add_meta_info(doc)
        
        # 收集数据并构建内容
        data = self.collect_data()
        self.build_content(doc, data)
        
        # 保存
        doc.save(str(filepath))
        return filepath
    
    def _get_filename_prefix(self) -> str:
        """获取文件名前缀（子类可覆盖）"""
        return self.REPORT_NAME.replace(" ", "_")
    
    def _setup_styles(self, doc: Document) -> None:
        """设置文档样式"""
        style = doc.styles['Normal']
        font = style.font
        font.name = '微软雅黑'
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    def _add_meta_info(self, doc: Document) -> None:
        """添加报告元信息"""
        from app.reports.utils import format_period_label, get_week_number
        
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        period_label = format_period_label(self.period)
        if "week" in self.period:
            period_label += f"（{get_week_number()}）"
        else:
            period_label += f"（{self.generated_at.strftime('%Y-%m-%d')}）"
        
        meta_text = (
            f"报告周期：{period_label}\n"
            f"生成时间：{self.generated_at.strftime('%Y-%m-%d %H:%M:%S')}"
        )
        run = meta.add_run(meta_text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        
        doc.add_paragraph()  # 空行分隔
