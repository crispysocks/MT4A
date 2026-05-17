"""学生心理健康周报生成器"""
from datetime import datetime
from typing import Optional
import os
import json

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlmodel import Session, select
from anthropic import Anthropic

from app.reports.base import ReportGenerator
from app.db.models import PsychologyProfile, PsychologyWarning, ExamSchedule
from app.db.connection import engine


class PsychologyWeeklyGenerator(ReportGenerator):
    """学生心理健康周报生成器"""
    
    REPORT_NAME = "学生心理健康周报"
    
    def __init__(self, period: str = "this_week", start_date: Optional[datetime] = None, end_date: Optional[datetime] = None):
        super().__init__(period, start_date, end_date)
        self.llm = Anthropic(
            base_url=os.getenv("LLM_BASE_URL"),
            api_key=os.getenv("LLM_AUTH_TOKEN"),
        )
        self.model = os.environ.get("MODEL_ID", "qwen3.6-plus")
    
    def _get_filename_prefix(self) -> str:
        from app.reports.utils import format_period_label
        period_label = format_period_label(self.period)
        date_str = self.generated_at.strftime("%Y%m%d")
        return f"心理健康周报_{period_label}_{date_str}"
    
    def collect_data(self) -> dict:
        """收集心理健康数据"""
        start_date = self.start_date
        end_date = self.end_date
        
        if start_date is None or end_date is None:
            from app.reports.utils import get_date_range
            start_date, end_date = get_date_range(self.period)
        
        with Session(engine) as session:
            # 情绪档案
            profiles = session.exec(
                select(PsychologyProfile).where(
                    PsychologyProfile.recorded_at >= start_date,
                    PsychologyProfile.recorded_at <= end_date
                )
            ).all()
            
            # 情绪标签统计
            emotion_counts = {}
            for p in profiles:
                tag = p.emotion_tag or "未记录"
                emotion_counts[tag] = emotion_counts.get(tag, 0) + 1
            
            # 心理预警
            warnings = session.exec(
                select(PsychologyWarning).where(
                    PsychologyWarning.created_at >= start_date,
                    PsychologyWarning.created_at <= end_date
                )
            ).all()
            
            # 考试安排
            exams = session.exec(
                select(ExamSchedule).where(
                    ExamSchedule.deadline >= start_date,
                    ExamSchedule.deadline <= end_date
                )
            ).all()
            
            return {
                "profiles": profiles,
                "emotion_counts": emotion_counts,
                "warnings": warnings,
                "exams": exams,
                "total_profiles": len(profiles),
                "total_warnings": len(warnings),
                "total_exams": len(exams),
                "start_date": start_date,
                "end_date": end_date,
            }
    
    def _ai_analyze(self, data: dict) -> str:
        """使用 LLM 生成心理态势分析"""
        prompt = f"""作为学生心理健康分析专家，请根据以下数据生成周报：

- 情绪打卡记录：{data['total_profiles']} 条
- 情绪分布：{json.dumps(data['emotion_counts'], ensure_ascii=False)}
- 心理预警数：{data['total_warnings']} 条
- 本周考试数：{data['total_exams']} 场

请分析：
1. 整体心理态势（情绪趋势、主要压力源）
2. 风险学生群体识别（孤独感、学业焦虑、文化冲突等）
3. 考试周/假期特殊节点影响
4. 个性化疏导建议与社群支持推荐

请直接输出分析内容。"""
        
        try:
            response = self.llm.messages.create(
                model=self.model,
                system="你是一个专业的学生心理健康分析专家。",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=2000,
            )
            
            for block in response.content:
                if hasattr(block, "text") and block.text:
                    return block.text.strip()
        except Exception as e:
            print(f"[AI 分析失败] {e}")
        
        return "AI 分析暂时不可用。"
    
    def build_content(self, doc: Document, data: dict) -> None:
        """构建报告内容"""
        # 1. 总体概况
        doc.add_heading("一、总体概况", level=1)
        overview = doc.add_paragraph()
        overview.add_run(f"情绪打卡记录：").bold = True
        overview.add_run(f"{data['total_profiles']} 条\n")
        overview.add_run(f"心理预警：").bold = True
        overview.add_run(f"{data['total_warnings']} 条\n")
        overview.add_run(f"本周考试：").bold = True
        overview.add_run(f"{data['total_exams']} 场")
        
        # 2. 情绪分布
        doc.add_heading("二、情绪分布", level=1)
        for emotion, count in data['emotion_counts'].items():
            doc.add_paragraph(f"{emotion}：{count} 人次")
        
        # 3. 心理预警明细
        if data['warnings']:
            doc.add_heading("三、心理预警明细", level=1)
            for w in data['warnings']:
                para = doc.add_paragraph()
                para.add_run(f"学生ID：{w.student_id} | 风险等级：{w.risk_level} | 状态：{w.status}")
                para.add_run(f"\n触发原因：{w.trigger_reason}")
        
        # 4. AI 心理态势分析
        doc.add_heading("四、AI 心理态势分析", level=1)
        analysis = self._ai_analyze(data)
        doc.add_paragraph(analysis)
