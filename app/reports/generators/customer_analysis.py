"""全域客户经营分析报告生成器"""
from datetime import datetime
from typing import Optional
import os
import json

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlmodel import Session, select
from anthropic import Anthropic

from app.reports.base import ReportGenerator
from app.db.models import Lead, LeadFollowUp, Registration
from app.db.connection import engine


class CustomerAnalysisGenerator(ReportGenerator):
    """全域客户经营分析报告生成器"""
    
    REPORT_NAME = "全域客户经营分析报告"
    
    def __init__(self, period: str = "this_week", start_date: Optional[datetime] = None, end_date: Optional[datetime] = None):
        super().__init__(period, start_date, end_date)
        self.llm = Anthropic(
            base_url=os.getenv("LLM_BASE_URL"),
            api_key=os.getenv("LLM_AUTH_TOKEN"),
        )
        self.model = os.environ.get("MODEL_ID", "qwen3.6-plus")
    
    def _get_filename_prefix(self) -> str:
        from app.reports.utils import format_period_label
        period_label = format_period_label(self.period)
        date_str = self.generated_at.strftime("%Y%m%d")
        return f"客户经营分析_{period_label}_{date_str}"
    
    def collect_data(self) -> dict:
        """收集客户数据"""
        start_date = self.start_date
        end_date = self.end_date
        
        if start_date is None or end_date is None:
            from app.reports.utils import get_date_range
            start_date, end_date = get_date_range(self.period)
        
        with Session(engine) as session:
            # 意向客户统计
            leads = session.exec(
                select(Lead).where(
                    Lead.created_at >= start_date,
                    Lead.created_at <= end_date
                )
            ).all()
            
            # 按状态分类
            status_counts = {}
            for lead in leads:
                status = lead.status or "未知"
                status_counts[status] = status_counts.get(status, 0) + 1
            
            # 跟进记录统计
            follow_ups = session.exec(
                select(LeadFollowUp).where(
                    LeadFollowUp.created_at >= start_date,
                    LeadFollowUp.created_at <= end_date
                )
            ).all()
            
            # 报名统计
            registrations = session.exec(
                select(Registration).where(
                    Registration.created_at >= start_date,
                    Registration.created_at <= end_date
                )
            ).all()
            
            return {
                "leads": leads,
                "status_counts": status_counts,
                "follow_ups": follow_ups,
                "registrations": registrations,
                "total_leads": len(leads),
                "total_follow_ups": len(follow_ups),
                "total_registrations": len(registrations),
                "start_date": start_date,
                "end_date": end_date,
            }
    
    def _ai_analyze(self, data: dict) -> str:
        """使用 LLM 生成客户经营洞察"""
        prompt = f"""作为客户经营分析专家，请根据以下数据生成经营分析报告：

- 新增意向客户：{data['total_leads']} 人
- 客户状态分布：{json.dumps(data['status_counts'], ensure_ascii=False)}
- 跟进记录数：{data['total_follow_ups']} 条
- 活动报名数：{data['total_registrations']} 人

请从以下维度分析：
1. 获客情况（新增趋势、渠道效果）
2. 转化情况（高价值客户特征）
3. 流失风险（流失原因、预警建议）
4. 经营建议（下一步行动方向）

请直接输出分析内容，不要使用 JSON 格式。"""
        
        try:
            response = self.llm.messages.create(
                model=self.model,
                system="你是一个专业的客户经营分析专家。",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=2000,
            )
            
            for block in response.content:
                if hasattr(block, "text") and block.text:
                    return block.text.strip()
        except Exception as e:
            print(f"[AI 分析失败] {e}")
        
        return "AI 分析暂时不可用。"
    
    def build_content(self, doc: Document, data: dict) -> None:
        """构建报告内容"""
        # 1. 总体概况
        doc.add_heading("一、总体概况", level=1)
        overview = doc.add_paragraph()
        overview.add_run(f"新增意向客户：").bold = True
        overview.add_run(f"{data['total_leads']} 人\n")
        overview.add_run(f"跟进记录：").bold = True
        overview.add_run(f"{data['total_follow_ups']} 条\n")
        overview.add_run(f"活动报名：").bold = True
        overview.add_run(f"{data['total_registrations']} 人")
        
        # 2. 客户状态分布
        doc.add_heading("二、客户状态分布", level=1)
        for status, count in data['status_counts'].items():
            doc.add_paragraph(f"{status}：{count} 人")
        
        # 3. AI 经营洞察
        doc.add_heading("三、AI 经营洞察", level=1)
        insight = self._ai_analyze(data)
        doc.add_paragraph(insight)
