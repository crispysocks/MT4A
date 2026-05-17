"""员工日报汇总报告生成器"""
from datetime import datetime
from typing import Optional
import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlmodel import Session, select
from anthropic import Anthropic

from app.reports.base import ReportGenerator
from app.db.models import DailyReport, User
from app.db.connection import engine


class DailyReportSummaryGenerator(ReportGenerator):
    """员工日报汇总报告生成器（支持日报/周报）"""
    
    REPORT_NAME = "员工日报汇总报告"
    
    def __init__(self, period: str = "today", start_date: Optional[datetime] = None, end_date: Optional[datetime] = None):
        super().__init__(period, start_date, end_date)
        self.llm = Anthropic(
            base_url=os.getenv("LLM_BASE_URL"),
            api_key=os.getenv("LLM_AUTH_TOKEN"),
        )
        self.model = os.environ.get("MODEL_ID", "qwen3.6-plus")
    
    def _get_filename_prefix(self) -> str:
        from app.reports.utils import format_period_label
        period_label = format_period_label(self.period)
        date_str = self.generated_at.strftime("%Y%m%d")
        return f"员工日报汇总_{period_label}_{date_str}"
    
    def collect_data(self) -> dict:
        """收集日报数据"""
        start_date = self.start_date
        end_date = self.end_date
        
        if start_date is None or end_date is None:
            from app.reports.utils import get_date_range
            start_date, end_date = get_date_range(self.period)
        
        with Session(engine) as session:
            # 查询指定时间范围内的所有日报
            reports = session.exec(
                select(DailyReport).where(
                    DailyReport.submitted_at >= start_date,
                    DailyReport.submitted_at <= end_date
                )
            ).all()
            
            # 获取关联的用户信息
            user_ids = list(set(r.user_id for r in reports))
            users = session.exec(
                select(User).where(User.id.in_(user_ids))
            ).all()
            user_map = {u.id: u.username for u in users}
            
            # 组装数据
            report_list = []
            for r in reports:
                report_list.append({
                    "user_id": r.user_id,
                    "username": user_map.get(r.user_id, "未知"),
                    "content": r.content,
                    "summary": r.summary,
                    "department": r.department,
                    "submitted_at": r.submitted_at,
                })
            
            return {
                "reports": report_list,
                "total_count": len(report_list),
                "start_date": start_date,
                "end_date": end_date,
            }
    
    def _ai_analyze(self, reports_text: str) -> dict:
        """使用 LLM 分析日报内容，提取核心进展、关键产出、潜在风险"""
        if not reports_text:
            return {
                "core_progress": "无数据",
                "key_outputs": "无数据",
                "risks": "无数据",
            }
        
        prompt = f"""请分析以下员工日报内容，提取并总结：
1. 核心进展（主要工作推进情况）
2. 关键产出（具体成果、数据指标）
3. 潜在风险（需要关注的问题、延期风险等）

日报内容：
{reports_text}

请以 JSON 格式返回，包含以下字段：
{{
  "core_progress": "...",
  "key_outputs": "...",
  "risks": "..."
}}"""
        
        try:
            response = self.llm.messages.create(
                model=self.model,
                system="你是一个专业的数据分析助手，擅长从文本中提取关键信息。",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=2000,
            )
            
            import json
            for block in response.content:
                if hasattr(block, "text") and block.text:
                    # 尝试解析 JSON
                    text = block.text.strip()
                    if text.startswith("```"):
                        # 去除 markdown 代码块
                        lines = text.split("\n")
                        text = "\n".join(lines[1:-1]) if len(lines) > 2 else text
                    
                    try:
                        return json.loads(text)
                    except json.JSONDecodeError:
                        pass
                    
                    return {
                        "core_progress": text,
                        "key_outputs": "",
                        "risks": "",
                    }
        except Exception as e:
            print(f"[AI 分析失败] {e}")
        
        return {
            "core_progress": "AI 分析失败",
            "key_outputs": "",
            "risks": "",
        }
    
    def build_content(self, doc: Document, data: dict) -> None:
        """构建报告内容"""
        reports = data["reports"]
        
        # 1. 总体概况
        doc.add_heading("一、总体概况", level=1)
        doc.add_paragraph(f"报告周期内共提交 {data['total_count']} 份日报。")
        
        if data["total_count"] == 0:
            doc.add_paragraph("本周期内无日报提交记录。")
            return
        
        # 2. AI 智能摘要
        doc.add_heading("二、AI 智能摘要", level=1)
        
        # 拼接所有日报内容用于 AI 分析
        all_content = "\n\n---\n\n".join(
            f"【{r['username']}】({r['submitted_at'].strftime('%Y-%m-%d %H:%M')})\n{r['content']}"
            for r in reports
        )
        
        analysis = self._ai_analyze(all_content)
        
        # 核心进展
        doc.add_heading("核心进展", level=2)
        doc.add_paragraph(analysis["core_progress"])
        
        # 关键产出
        doc.add_heading("关键产出", level=2)
        doc.add_paragraph(analysis["key_outputs"])
        
        # 潜在风险
        doc.add_heading("潜在风险", level=2)
        risk_para = doc.add_paragraph(analysis["risks"])
        if analysis["risks"] and analysis["risks"] != "无数据":
            for run in risk_para.runs:
                run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        
        # 3. 明细记录
        doc.add_heading("三、日报明细", level=1)
        
        for r in reports:
            # 员工信息
            header = doc.add_paragraph()
            run = header.add_run(f"【{r['username']}】")
            run.bold = True
            run.font.size = Pt(12)
            
            if r.get("department"):
                header.add_run(f" | 部门：{r['department']}")
            
            time_run = header.add_run(f" | 提交时间：{r['submitted_at'].strftime('%Y-%m-%d %H:%M')}")
            time_run.font.size = Pt(9)
            time_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            
            # 日报内容
            doc.add_paragraph(r["content"])
            
            # 如果有摘要，显示摘要
            if r.get("summary"):
                summary_para = doc.add_paragraph()
                summary_run = summary_para.add_run(f"摘要：{r['summary']}")
                summary_run.font.size = Pt(10)
                summary_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
                summary_run.italic = True
            
            doc.add_paragraph()  # 分隔空行
